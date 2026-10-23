"""ASSISTANT-F19 (ASSISTANT-G55) · Read what a question attaches: an image goes to the model as it is (a model that
cannot see gets a note instead), a document (PDF, Word, Excel, text, CSV, Markdown, JSON) is turned into its text.

The Hub sends each file inside the message, in base64: erplora.com downloads nothing (no address of a document or an
image is fetched, so there is no door to the internal network) and writes nothing to disk. Each attachment has a
size limit (``MAX_ATTACHMENT_BYTES`` of base64, about 9 MB; over it the turn is refused with 413, never cut) and a
conversation a number of them; their data does not count in the 1 MB of text. What cannot be read becomes a note for
the model, never a failed turn: the old product's ``internal_error`` with a corrupt PDF, ``.docx`` or ``.xlsx``
(SAAS_ASSISTANT-F18, «parcial») is a test here. Cases of the old ``test_document_ingest.py`` and
``test_stateless_bridge.py::TestPayloadCap``; the file's name never reaches a log (saas#2445).
"""

import base64
import io
import logging
import zipfile

import pytest
from django.test import Client

from assistant import attachments
from assistant.tests.conftest import openai_answer
from assistant.tests.doors import door, error_code, files_written, frames
from assistant.turn import TurnRefused, prepare

PNG = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
COMPLETIONS = "/v1/chat/completions"


def b64(raw: bytes) -> str:
    return base64.b64encode(raw).decode()


def pdf_with(text: str) -> bytes:
    from reportlab.pdfgen import canvas

    out = io.BytesIO()
    page = canvas.Canvas(out)
    page.drawString(72, 720, text)
    page.save()
    return out.getvalue()


def blank_pdf() -> bytes:
    from reportlab.pdfgen import canvas

    out = io.BytesIO()
    page = canvas.Canvas(out)
    page.showPage()
    page.save()
    return out.getvalue()


def docx_with(*paragraphs: str) -> bytes:
    from docx import Document

    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def xlsx_with(sheets: dict) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def document(data: bytes, mime: str, name: str) -> dict:
    return {"type": "input_file", "filename": name, "mime_type": mime, "data": b64(data)}


def asking(*parts) -> list[dict]:
    return [{"role": "user", "content": [{"type": "text", "text": "Read this"}, *parts]}]


def expanded_parts(*parts) -> list[dict]:
    return prepare(asking(*parts), "")[1]["content"]


class TestDocumentsBecomeText:
    def test_a_pdf_becomes_its_text(self):
        part = expanded_parts(document(pdf_with("Carta de vinos 2026"), "application/pdf", "carta.pdf"))[1]

        assert part["type"] == "text"
        assert part["text"].startswith("[Attachment: carta.pdf]")
        assert "Carta de vinos 2026" in part["text"]

    def test_a_word_document_becomes_its_paragraphs(self):
        data = docx_with("Menú del día", "Lentejas")
        name = "menu.docx"
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        text = expanded_parts(document(data, mime, name))[1]["text"]

        assert "Menú del día\nLentejas" in text

    def test_a_spreadsheet_becomes_a_section_per_sheet_without_empty_rows(self):
        data = xlsx_with({"Stock": [["Producto", "Cantidad"], [None, None], ["Café", 12]], "Precios": [["Café", 1.5]]})
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        text = expanded_parts(document(data, mime, "stock.xlsx"))[1]["text"]

        assert "# Stock\nProducto | Cantidad\nCafé | 12" in text
        assert "# Precios\nCafé | 1.5" in text

    @pytest.mark.parametrize(("mime", "name"), [("text/plain", "notas.txt"), ("text/csv", "ventas.csv"), ("", "readme.md"), ("application/json", "datos.json")])
    def test_plain_text_kinds_are_read_as_they_are(self, mime, name):
        text = expanded_parts(document("línea uno\nlínea dos".encode(), mime, name))[1]["text"]

        assert "línea uno\nlínea dos" in text

    def test_a_long_document_is_cut_to_what_the_model_can_read(self):
        text = expanded_parts(document(("x" * (attachments.MAX_EXTRACTED_CHARS + 50)).encode(), "text/plain", "long.txt"))[1]["text"]

        assert text.endswith("[…cut…]")
        assert len(text) < attachments.MAX_EXTRACTED_CHARS + 100

    def test_a_scanned_pdf_says_it_has_no_text(self):
        text = expanded_parts(document(blank_pdf(), "application/pdf", "escaneo.pdf"))[1]["text"]

        assert "no text" in text and "escaneo.pdf" in text

    def test_a_data_uri_document_is_read_too(self):
        part = {"type": "input_file", "filename": "n.txt", "mime_type": "text/plain", "data": "data:text/plain;base64," + b64(b"hola")}

        assert "hola" in expanded_parts(part)[1]["text"]

    def test_nothing_is_written_to_disk(self):
        pdf, docx, xlsx = pdf_with("pdf"), docx_with("doc"), xlsx_with({"S": [["a"]]})
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        with files_written() as written:
            parts = expanded_parts(
                document(pdf, "application/pdf", "a.pdf"),
                document(docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "b.docx"),
                document(xlsx, mime, "c.xlsx"),
            )

        assert written == []
        assert [("pdf" in parts[1]["text"]), ("doc" in parts[2]["text"]), ("a" in parts[3]["text"])] == [True, True, True]

    def test_the_watch_sees_a_file_written_by_a_library(self):
        """The control catches the positive: building a workbook makes openpyxl write a temporary file."""
        with files_written() as written:
            xlsx_with({"S": [["a"]]})

        assert written


class TestWhatCannotBeReadIsANote:
    @pytest.mark.parametrize(
        ("data", "mime", "name"),
        [
            (b"%PDF-1.4 broken", "application/pdf", "rota.pdf"),
            (b"PK\x03\x04 not a zip", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "rota.docx"),
            (b"PK\x03\x04 not a zip", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "rota.xlsx"),
        ],
    )
    def test_a_corrupt_document_is_a_note_not_a_failed_turn(self, data, mime, name):
        text = expanded_parts(document(data, mime, name))[1]["text"]

        assert text.startswith(f"[Attachment {name}: could not be read")

    def test_an_old_word_document_is_a_note(self):
        text = expanded_parts(document(b"\xd0\xcf\x11\xe0", "application/msword", "viejo.doc"))[1]["text"]

        assert "could not be read" in text

    def test_broken_base64_is_a_note(self):
        part = {"type": "input_file", "filename": "x.txt", "mime_type": "text/plain", "data": "@@not base64@@"}

        assert "could not be read" in expanded_parts(part)[1]["text"]

    def test_a_document_by_address_is_not_downloaded(self, fake_provider):
        part = {"type": "input_file", "filename": "x.pdf", "file_url": f"{fake_provider.url}/x.pdf"}

        assert "could not be read" in expanded_parts(part)[1]["text"]
        assert fake_provider.requests == []

    def test_a_zip_bomb_is_a_note(self):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"0" * (attachments.MAX_UNZIPPED_BYTES + 1))
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        assert "unfolds into too much" in expanded_parts(document(out.getvalue(), mime, "bomb.docx"))[1]["text"]

    def test_the_name_of_an_unreadable_file_never_reaches_the_log(self, caplog):
        caplog.set_level(logging.DEBUG)

        expanded_parts(document(b"%PDF broken", "application/pdf", "Nómina de Rosalía.pdf"))
        expanded_parts(document(b"???", "", "Contrato Rosalía.doc"))

        logged = "\n".join(record.getMessage() for record in caplog.records)
        assert "Rosalía" not in logged


class TestImages:
    def image(self, url: str) -> dict:
        return {"type": "image_url", "image_url": {"url": url}}

    def test_an_image_goes_to_a_model_that_sees_as_it_is(self):
        part = self.image(f"data:image/png;base64,{PNG}")

        assert attachments.for_model(asking(part), vision=True)[0]["content"][1] == part

    def test_a_model_that_cannot_see_gets_a_note_instead(self):
        content = attachments.for_model(asking(self.image(f"data:image/png;base64,{PNG}")), vision=False)[0]["content"]

        assert content[1]["type"] == "text"
        assert "cannot see images" in content[1]["text"]

    @pytest.mark.parametrize("url", ["https://example.com/a.png", "http://169.254.169.254/latest", "data:text/html;base64,PGI+"])
    def test_an_image_by_address_or_not_an_image_is_a_note(self, url):
        content = prepare(asking(self.image(url)), "")[1]["content"]

        assert content[1]["type"] == "text"
        assert "could not be read" in content[1]["text"]


class TestLimits:
    def test_an_attachment_over_its_limit_refuses_the_turn(self):
        big = "A" * (attachments.MAX_ATTACHMENT_BYTES + 4)

        with pytest.raises(TurnRefused) as refused:
            prepare(asking({"type": "input_file", "filename": "big.txt", "mime_type": "text/plain", "data": big}), "")

        assert (refused.value.status, refused.value.code) == (413, "payload_too_large")

    def test_too_many_attachments_refuse_the_turn(self):
        parts = [{"type": "image_url", "image_url": {"url": f"data:image/png;base64,{PNG}"}}] * (attachments.MAX_ATTACHMENTS + 1)

        with pytest.raises(TurnRefused) as refused:
            prepare(asking(*parts), "")

        assert refused.value.status == 413

    def test_attachments_do_not_count_in_the_text(self):
        almost = "A" * 900_000
        parts = [{"type": "input_file", "filename": f"{n}.txt", "mime_type": "text/plain", "data": almost} for n in range(2)]

        prepare(asking(*parts), "")

    def test_the_text_is_still_limited_with_an_attachment(self):
        with pytest.raises(TurnRefused):
            prepare([{"role": "user", "content": [{"type": "text", "text": "x" * 1_000_001}, {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{PNG}"}}]}], "")


@pytest.mark.django_db
class TestInTheTurn:
    def test_the_model_reads_the_documents_text_and_sees_the_image(self, internal, business, answering, fake_provider):
        fake_provider.reply("POST", COMPLETIONS, sse=openai_answer("Ok"))
        conversation = asking(
            document(pdf_with("Factura 42"), "application/pdf", "factura.pdf"),
            {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{PNG}"}},
        )

        frames(door(Client(), "/internal/assistant/chat/stream", {"hub": business.slug, "messages": conversation}))

        [sent] = [seen for seen in fake_provider.requests if seen.path.endswith("/chat/completions")]
        parts = sent.body["messages"][1]["content"]
        assert "Factura 42" in parts[1]["text"]
        assert parts[2]["type"] == "image_url"

    def test_an_attachment_too_big_is_a_413_before_the_model(self, internal, business, answering, fake_provider):
        big = {"type": "input_file", "filename": "big.txt", "mime_type": "text/plain", "data": "A" * (attachments.MAX_ATTACHMENT_BYTES + 4)}

        response = door(Client(), "/internal/assistant/chat", {"hub": business.slug, "messages": asking(big)})

        assert (response.status_code, error_code(response)) == (413, "payload_too_large")
        assert fake_provider.requests == []
