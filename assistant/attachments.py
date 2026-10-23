"""The attachments of a question (ASSISTANT-G55): images for the model to see, documents turned into their text.

A user message's content may be a list of parts. Besides ``{"type": "text"}``:

- ``{"type": "image_url", "image_url": {"url": "data:image/<png|jpeg|webp|gif>;base64,…"}}``: goes to the model as it
  is if it sees images (``AIModel.supports_vision``); otherwise, or if it is not such an image (an address, another
  type), a note in its place. erplora.com downloads no image: a provider would fetch an address on its own.
- ``{"type": "input_file", "filename", "mime_type", "data": "<base64>"}``: a PDF (pypdf), a Word document
  (python-docx), a spreadsheet (openpyxl, a section per sheet, cells joined by « | »), or plain text, CSV, Markdown or
  JSON, turned into a ``text`` part ``[Attachment: <name>]`` with at most ``MAX_EXTRACTED_CHARS``. A document by
  address (``file_url``) is not downloaded.

Everything happens in memory: nothing is written to disk. A Word document or a spreadsheet is a zip, checked before
it is opened (``MAX_UNZIPPED_BYTES``, ``MAX_ZIP_ENTRIES``), so a small file cannot unfold into gigabytes. What cannot
be read becomes a note for the model, never a failed turn; the file's name stays out of the logs (saas#2445). Each
attachment's data has a limit (``MAX_ATTACHMENT_BYTES``) and a conversation a number of them (``MAX_ATTACHMENTS``):
over them the turn is refused (``turn.TurnRefused``, 413), never cut; their data does not count in the text's limit.
"""

from __future__ import annotations

import base64
import binascii
import io
import logging
import os
import zipfile

logger = logging.getLogger(__name__)

#: The base64 of one attachment (about 9 MB of file).
MAX_ATTACHMENT_BYTES = 12 * 1024 * 1024
#: The attachments of a conversation (the Hub sends it whole on every turn).
MAX_ATTACHMENTS = 20
#: The text of one document the model reads.
MAX_EXTRACTED_CHARS = 100_000
#: The pages of a PDF that are read.
MAX_PDF_PAGES = 300
#: What a Word document or a spreadsheet may unfold into, and its files.
MAX_UNZIPPED_BYTES = 64 * 1024 * 1024
MAX_ZIP_ENTRIES = 10_000
IMAGE_TYPES = frozenset({"image/png", "image/jpeg", "image/webp", "image/gif"})
WORD = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SHEET = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TEXT_EXTENSIONS = (".txt", ".md", ".csv", ".json")


class Unreadable(Exception):
    """A document whose text cannot be read; the message says why (never its name)."""


def parts(messages: list[dict]):
    """``(message, index, part)`` of every part of a list content."""
    for message in messages:
        content = message.get("content")
        if isinstance(content, list):
            for index, part in enumerate(content):
                if isinstance(part, dict):
                    yield message, index, part


def blob(part: dict) -> str | None:
    """The base64 an attachment part carries, or ``None``."""
    if part.get("type") == "image_url":
        url = (part.get("image_url") or {}).get("url") if isinstance(part.get("image_url"), dict) else None
        return url if isinstance(url, str) and url.startswith("data:") else None
    if part.get("type") == "input_file":
        data = part.get("data")
        return data if isinstance(data, str) and data else None
    return None


def check(messages: list[dict]) -> None:
    """Refuse (413) a conversation with an attachment over ``MAX_ATTACHMENT_BYTES`` or more than ``MAX_ATTACHMENTS``."""
    from assistant.turn import TurnRefused

    count = 0
    for _message, _index, part in parts(messages):
        if part.get("type") not in ("image_url", "input_file"):
            continue
        count += 1
        data = blob(part)
        if data is not None and len(data) > MAX_ATTACHMENT_BYTES:
            raise TurnRefused(413, "payload_too_large", f"An attachment exceeds {MAX_ATTACHMENT_BYTES} bytes of base64.")
    if count > MAX_ATTACHMENTS:
        raise TurnRefused(413, "payload_too_large", f"The conversation has more than {MAX_ATTACHMENTS} attachments.")


def without_data(messages: list[dict]) -> list[dict]:
    """``messages`` with each attachment's data left out, to measure the text."""
    measured = []
    for message in messages:
        content = message.get("content")
        if not isinstance(content, list):
            measured.append(message)
            continue
        measured.append({**message, "content": [_dataless(part) for part in content]})
    return measured


def _dataless(part):
    if not isinstance(part, dict) or blob(part) is None:
        return part
    if part.get("type") == "image_url":
        return {**part, "image_url": {"url": "<image>"}}
    return {**part, "data": "<data>"}


def expand(messages: list[dict]) -> list[dict]:
    """``messages`` with each document turned into its text and each image that is not one the model may see turned
    into a note. Never raises."""
    expanded = []
    for message in messages:
        content = message.get("content")
        if not isinstance(content, list):
            expanded.append(message)
            continue
        expanded.append({**message, "content": [_expanded(part) for part in content]})
    return expanded


def _expanded(part):
    if not isinstance(part, dict):
        return part
    if part.get("type") == "input_file":
        return _document(part)
    if part.get("type") == "image_url":
        url = blob(part)
        if url is None or _data_type(url) not in IMAGE_TYPES:
            return note("image", "only an image inside the message (PNG, JPEG, WebP or GIF) is read")
    return part


def for_model(messages: list[dict], *, vision: bool) -> list[dict]:
    """The conversation as a model sees it: a model that cannot see images gets a note in each one's place."""
    if vision:
        return messages
    seen = []
    for message in messages:
        content = message.get("content")
        if not isinstance(content, list):
            seen.append(message)
            continue
        seen.append(
            {
                **message,
                "content": [
                    {"type": "text", "text": "[An image: this model cannot see images; tell the person so.]"}
                    if isinstance(part, dict) and part.get("type") == "image_url"
                    else part
                    for part in content
                ],
            }
        )
    return seen


def note(name: str, reason: str) -> dict:
    return {"type": "text", "text": f"[Attachment {name}: could not be read ({reason}).]"}


def _data_type(url: str) -> str:
    head = url[len("data:") : url.find(",")] if "," in url else ""
    return head.split(";")[0].strip().lower()


def _decoded(data: str) -> bytes:
    text = data.strip()
    if text.startswith("data:"):
        if "," not in text:
            raise Unreadable("malformed data address")
        text = text[text.find(",") + 1 :]
    try:
        return base64.b64decode(text, validate=True)
    except (binascii.Error, ValueError):
        raise Unreadable("not valid base64") from None


def _document(part: dict) -> dict:
    name = str(part.get("filename") or part.get("name") or "document")
    mime = str(part.get("mime_type") or "").lower()
    data = part.get("data")
    if not isinstance(data, str) or not data:
        reason = "documents are read only inside the message" if part.get("file_url") or part.get("url") else "empty"
        return note(name, reason)
    try:
        text = extract(_decoded(data), mime, name)
    except Unreadable as problem:
        logger.info("assistant.attachments.unreadable type=%s reason=%s", _kind(mime, name), problem)
        return note(name, str(problem))
    except Exception as problem:  # noqa: BLE001 - a broken file of any library is a note, never a failed turn
        logger.info("assistant.attachments.unreadable type=%s reason=%s", _kind(mime, name), type(problem).__name__)
        return note(name, "the file is damaged")
    text = text.strip()
    if not text:
        return {
            "type": "text",
            "text": f"[Attachment {name}: no text in it (likely a scanned document); if asked about it, say you could not read it.]",
        }
    if len(text) > MAX_EXTRACTED_CHARS:
        text = text[:MAX_EXTRACTED_CHARS] + "\n[…cut…]"
    return {"type": "text", "text": f"[Attachment: {name}]\n{text}"}


def _kind(mime: str, name: str) -> str:
    """What a log may say of a file: its type or its extension, never its name."""
    return mime or os.path.splitext(name)[1].lower() or "unknown"


def extract(raw: bytes, mime: str, name: str) -> str:
    lowered = name.lower()
    if mime == "application/pdf" or lowered.endswith(".pdf"):
        return _pdf(raw)
    if mime == WORD or lowered.endswith(".docx"):
        return _docx(raw)
    if mime == SHEET or lowered.endswith(".xlsx"):
        return _xlsx(raw)
    if mime.startswith("text/") or mime == "application/json" or lowered.endswith(TEXT_EXTENSIONS):
        return raw.decode("utf-8", errors="replace")
    raise Unreadable("this kind of file is not read: send a PDF, Word (.docx), Excel (.xlsx) or text file")


def _pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    pages = []
    total = 0
    for page in reader.pages[:MAX_PDF_PAGES]:
        text = page.extract_text() or ""
        pages.append(text)
        total += len(text)
        if total > MAX_EXTRACTED_CHARS:
            break
    return "\n".join(pages)


def _unzippable(raw: bytes) -> None:
    """A zip that unfolds into no more than ``MAX_UNZIPPED_BYTES`` in ``MAX_ZIP_ENTRIES`` files, or ``Unreadable``."""
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile:
        raise Unreadable("the file is damaged") from None
    if len(entries) > MAX_ZIP_ENTRIES or sum(entry.file_size for entry in entries) > MAX_UNZIPPED_BYTES:
        raise Unreadable("the file unfolds into too much")


def _docx(raw: bytes) -> str:
    from docx import Document

    _unzippable(raw)
    return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(raw)).paragraphs)


def _xlsx(raw: bytes) -> str:
    from openpyxl import load_workbook

    _unzippable(raw)
    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        sheets, total = [], 0
        for sheet in book.worksheets:
            lines = []
            for row in sheet.iter_rows(values_only=True):
                cells = list(row)
                while cells and cells[-1] is None:
                    cells.pop()
                if not cells:
                    continue
                line = " | ".join("" if cell is None else str(cell) for cell in cells)
                lines.append(line)
                total += len(line) + 1
                if total > MAX_EXTRACTED_CHARS:
                    break
            if lines:
                sheets.append(f"# {sheet.title}\n" + "\n".join(lines))
            if total > MAX_EXTRACTED_CHARS:
                break
        return "\n\n".join(sheets)
    finally:
        book.close()
